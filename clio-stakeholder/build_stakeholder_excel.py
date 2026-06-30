#!/usr/bin/env python3
"""
Clio Intressentanalys — Excel-mall v1.0
Bygger en återanvändbar Excel-template för intressentkartläggning (stakeholder
mapping) utifrån Slack-metadata. Filen innehåller ingen kunddata — rådatan
fyller du från en CSV/Analytics-export som du kör själv hos kunden.

Mallen har två lager:
  • Rådata  — metadata du läser ut ur Slack (aktivitet, roller, kanaler)
  • Analys  — Inflytande/Intresse → Mendelow-ruta (auto) + RACI och kommplan

Användning:
    python build_stakeholder_excel.py
    python build_stakeholder_excel.py --out ~/Intressentanalys.xlsx
    python build_stakeholder_excel.py --csv slack_export.csv --rows 60 --kund "Kund AB"
"""

from __future__ import annotations

import argparse
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

# ── AIAB Färgpalett ──────────────────────────────────────────
C_CREAM       = "F7F2E8"
C_CREAM_DARK  = "EDE5D0"
C_INPUT_BG    = "FDF6DC"
C_INPUT_BORD  = "C8A84B"
C_HEADER_DARK = "2A3F6F"
C_HEADER_MED  = "4A6FA5"
C_GOLD        = "C8A84B"
C_GOLD_LIGHT  = "E8C96A"
C_INK         = "3D2E0A"
C_INK_SOFT    = "7A6A4A"
C_FORMULA     = "2A3F6F"
C_INPUT_VAL   = "B8860B"
C_WHITE       = "FFFFFF"

# Mendelow-rutornas färger (makt/intresse-matrisen)
C_MANAGE  = "C9A0DC"   # Hantera nära     — lila
C_SATISFY = "F4C77B"   # Håll nöjd        — guld
C_INFORM  = "9BC4E2"   # Håll informerad  — blå
C_MONITOR = "CFD8C5"   # Bevaka           — grågrön

# Referens till Rådata-fliken i formler (citerad pga å)
RAD = "'Rådata'"


def fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def font(bold=False, italic=False, color=C_INK, size=10, name="Arial"):
    return Font(name=name, bold=bold, italic=italic, color=color, size=size)


def border_thin():
    s = Side(style="thin", color=C_GOLD)
    return Border(top=s, bottom=s, left=s, right=s)


def input_border():
    s = Side(style="thin", color=C_INPUT_BORD)
    return Border(top=s, bottom=s, left=s, right=s)


def _banner(ws, span, text, subtitle=None):
    """Sätter titelrad (och valfri underrubrik) överst på en flik."""
    ws.merge_cells(f"A1:{span}1")
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", bold=True, size=15, color=C_GOLD_LIGHT)
    ws["A1"].fill = fill(C_HEADER_DARK)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    if subtitle is not None:
        ws.merge_cells(f"A2:{span}2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(name="Arial", italic=True, size=9, color=C_INK_SOFT)
        ws["A2"].fill = fill(C_CREAM_DARK)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20


def _table_header(ws, row, headers, bg=C_HEADER_MED):
    for col, txt in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=txt)
        c.font = Font(name="Arial", bold=True, size=9, color=C_WHITE)
        c.fill = fill(bg)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = border_thin()
    ws.row_dimensions[row].height = 26


# ════════════════════════════════════════════════════════════
# FLIK 1 – Instructions
# ════════════════════════════════════════════════════════════
def build_instructions(wb):
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_properties.tabColor = C_HEADER_DARK
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 80

    _banner(
        ws, "C",
        "🧭 Clio Intressentanalys — Excel-mall v1.0",
        "Intressentkartläggning från Slack-metadata. Du äger rådatan; mallen bygger kartan.",
    )

    sections = [
        ("", "", ""),
        ("SYFTE", "", ""),
        ("", "Vad är det här?",
         "En återanvändbar mall för att kartlägga intressenter (stakeholders) i ett "
         "Slack-baserat kunduppdrag. Filen innehåller ingen kunddata — du fyller "
         "Rådata-fliken från en export du kör själv, och Analys-fliken räknar fram "
         "varje intressents position i Mendelow-matrisen automatiskt."),
        ("", "Princip",
         "Du gör datajobbet och behåller rådatan hos dig (integritet + konsultetik). "
         "Mallen behöver aldrig se kundens Slack — den strukturerar bara metadatan."),
        ("", "", ""),
        ("KOM IGÅNG", "", ""),
        ("", "Steg 1 — Exportera ur Slack",
         "Be en workspace-admin köra en Analytics-export (Slack-admin → Analytics) "
         "som ger medlemsaktivitet per kanal som CSV — UTAN meddelandeinnehåll. "
         "Komplettera med kanalägare, admins och User Groups."),
        ("", "Steg 2 — Fyll Rådata",
         "Klistra in/skriv en rad per intressent i fliken [Rådata]. Anonymisera det "
         "som är känsligt innan det lämnar kunden. Tips: starta build-scriptet med "
         "--csv din_export.csv för att förfylla automatiskt."),
        ("", "Steg 3 — Sätt Inflytande & Intresse",
         "I fliken [Analys], ge varje intressent Inflytande (1–5) och Intresse (1–5). "
         "Härled värdena ur metadatan — se tolkningstipsen nedan."),
        ("", "Steg 4 — Läs kartan",
         "Mendelow-rutan fylls i automatiskt. Fyll på med RACI och kommunikationsbehov. "
         "Summeringen längst ner i Analys räknar intressenter per ruta."),
        ("", "", ""),
        ("FLIKAR", "", ""),
        ("", "Instructions", "Den här sidan. Syfte, steg-för-steg och tolkningstips."),
        ("", "Settings", "Justerbara parametrar: kund, period, trösklar för Mendelow."),
        ("", "Rådata", "Metadata ur Slack. En rad per intressent. Inget meddelandeinnehåll."),
        ("", "Analys", "Inflytande/Intresse → Mendelow-ruta (auto), RACI och kommplan."),
        ("", "File History", "Versionslogg för filen."),
        ("", "", ""),
        ("MENDELOW (makt/intresse)", "", ""),
        ("", "Hantera nära", "Högt inflytande + högt intresse. Involvera tätt, förankra beslut."),
        ("", "Håll nöjd", "Högt inflytande + lågt intresse. Formell makt — håll nöjd, tråka inte ut."),
        ("", "Håll informerad", "Lågt inflytande + högt intresse. Ambassadörer — mata med info."),
        ("", "Bevaka", "Lågt inflytande + lågt intresse. Minsta möjliga insats, håll koll."),
        ("", "Trösklar", "Gränsen mellan högt/lågt sätts i [Settings] (default 3 = värde ≥ 3 är högt)."),
        ("", "", ""),
        ("RACI (per arbetsström)", "", ""),
        ("", "R / A / C / I",
         "Responsible (gör jobbet) / Accountable (äger beslutet) / Consulted (rådfrågas) "
         "/ Informed (informeras). Sätt per intressent och arbetsström."),
        ("", "", ""),
        ("TOLKNINGSTIPS", "", ""),
        ("", "Aktiv ≠ mäktig",
         "Den mest aktiva personen är sällan den mäktigaste — hög volym betyder ofta "
         "operativt nav, inte beslutsmandat."),
        ("", "Leta asymmetrin",
         "Någon som TAGGAS ofta men POSTAR sällan sitter oftast på beslut. Hög taggning "
         "+ låg egen aktivitet = ofta formell makt."),
        ("", "Tysta admins",
         "Glöm inte admin-konton som äger kanaler utan att synas i flödet — lätta att "
         "missa men formellt centrala."),
        ("", "Externa intressenter",
         "Gästkonton och Slack Connect-kanaler avslöjar leverantörer, kunder och "
         "partners. Markera dem som Gäst/Connect-partner i Rådata."),
        ("", "", ""),
        ("INTEGRITET", "", ""),
        ("", "Ingen live-koppling",
         "Mallen hämtar ingenting automatiskt. Du exporterar och fyller manuellt."),
        ("", "Du äger rådatan",
         "Anonymisera/ta bort känsligt innan det lämnar kunden. Mallen behöver bara "
         "den avskalade metadatan."),
    ]

    row = 3
    for label, key, val in sections:
        if label and not key:
            ws.merge_cells(f"A{row}:C{row}")
            ws[f"A{row}"] = f"  {label}"
            ws[f"A{row}"].font = Font(name="Arial", bold=True, size=10, color=C_WHITE)
            ws[f"A{row}"].fill = fill(C_HEADER_MED)
            ws.row_dimensions[row].height = 20
        elif key:
            ws[f"B{row}"] = key
            ws[f"B{row}"].font = Font(name="Arial", bold=True, size=9, color=C_INK)
            ws[f"B{row}"].fill = fill(C_CREAM)
            ws[f"B{row}"].alignment = Alignment(vertical="top", wrap_text=True)
            ws[f"C{row}"] = val
            ws[f"C{row}"].font = Font(name="Arial", size=9, color=C_INK_SOFT)
            ws[f"C{row}"].fill = fill(C_CREAM)
            ws[f"C{row}"].alignment = Alignment(wrap_text=True, vertical="top")
            est = max(1, (len(val) // 78) + val.count("\n") + 1)
            ws.row_dimensions[row].height = 14 * est + 4
        else:
            ws.row_dimensions[row].height = 8
        row += 1


# ════════════════════════════════════════════════════════════
# FLIK 2 – Settings
# ════════════════════════════════════════════════════════════
def build_settings(wb, kund):
    ws = wb.create_sheet("Settings")
    ws.sheet_properties.tabColor = C_GOLD
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 56

    _banner(
        ws, "D", "⚙️ Inställningar",
        "Gula celler = inmatningsfält. Trösklarna styr Mendelow-rutan i Analys.",
    )
    _table_header(ws, 3, ["", "Parameter", "Värde", "Förklaring"])

    rows = [
        ("Kund / projekt", kund, "Namn på kund eller uppdrag. Visas i Analys-rubriken."),
        ("Slack workspace", "[workspace.slack.com]", "Vilken workspace analysen gäller."),
        ("Analysperiod", "[t.ex. senaste 90 dagar]", "Perioden Analytics-exporten täcker."),
        ("Exportdatum", date.today().strftime("%Y-%m-%d"), "Datum då metadatan exporterades."),
        ("Tröskel inflytande (1–5)", 3, "Värde ≥ detta räknas som HÖGT inflytande i Mendelow."),
        ("Tröskel intresse (1–5)", 3, "Värde ≥ detta räknas som HÖGT intresse i Mendelow."),
        ("Analytiker", "[Ditt namn]", "Vem som gjort kartläggningen."),
    ]
    for i, (param, val, expl) in enumerate(rows):
        r = 4 + i
        ws[f"B{r}"] = param
        ws[f"B{r}"].font = Font(name="Arial", bold=True, size=9, color=C_INK)
        ws[f"B{r}"].fill = fill(C_CREAM)
        ws[f"B{r}"].border = border_thin()
        ws[f"C{r}"] = val
        ws[f"C{r}"].font = Font(name="Arial", bold=True, size=9, color=C_INPUT_VAL)
        ws[f"C{r}"].fill = fill(C_INPUT_BG)
        ws[f"C{r}"].border = input_border()
        ws[f"C{r}"].alignment = Alignment(horizontal="center")
        ws[f"D{r}"] = expl
        ws[f"D{r}"].font = Font(name="Arial", italic=True, size=9, color=C_INK_SOFT)
        ws[f"D{r}"].fill = fill(C_CREAM)
        ws[f"D{r}"].border = border_thin()
        ws[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 22

    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C8")
    dv.add("C9")


# ════════════════════════════════════════════════════════════
# FLIK 3 – Rådata
# ════════════════════════════════════════════════════════════
RADATA_HEADERS = [
    "#", "Namn", "Roll/titel", "Intern/Extern", "Admin/Ägare", "Nyckelkanaler",
    "Meddelanden (period)", "Aktiva dagar", "Ggr @-nämnd", "Svarar/initierar",
    "User groups",
]


def build_radata(wb, rows, csv_rows):
    ws = wb.create_sheet("Rådata")
    ws.sheet_properties.tabColor = C_HEADER_MED
    widths = [4, 24, 24, 14, 13, 30, 16, 13, 14, 22, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    _banner(ws, "K", "👥 Rådata — Slack-metadata (en rad per intressent)")
    _table_header(ws, 2, RADATA_HEADERS, bg=C_HEADER_DARK)

    data_start, data_end = 3, 2 + rows
    for idx in range(rows):
        r = data_start + idx
        bg = C_CREAM if idx % 2 == 0 else C_CREAM_DARK
        source = csv_rows[idx] if idx < len(csv_rows) else None
        for col in range(1, len(RADATA_HEADERS) + 1):
            c = ws.cell(row=r, column=col)
            if col == 1:
                c.value = idx + 1 if source else None
                c.font = Font(name="Arial", bold=True, size=9, color=C_INK_SOFT)
            else:
                if source and (col - 2) < len(source):
                    c.value = source[col - 2]
                c.font = Font(name="Arial", size=9, color=C_INK)
            c.fill = fill(bg)
            c.border = border_thin()
            c.alignment = Alignment(vertical="center", wrap_text=(col in (6, 11)))
        ws.row_dimensions[r].height = 18

    dv_typ = DataValidation(type="list", formula1='"Intern,Gäst,Connect-partner"',
                            allow_blank=True)
    dv_ja = DataValidation(type="list", formula1='"Ja,Nej"', allow_blank=True)
    dv_roll = DataValidation(
        type="list",
        formula1='"Svarar mest,Initierar mest,Taggad men tyst,Blandat"',
        allow_blank=True,
    )
    for dv, col in ((dv_typ, "D"), (dv_ja, "E"), (dv_roll, "J")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{data_start}:{col}{data_end}")

    ws.freeze_panes = "C3"


# ════════════════════════════════════════════════════════════
# FLIK 4 – Analys
# ════════════════════════════════════════════════════════════
ANALYS_HEADERS = [
    "#", "Namn", "Inflytande (1–5)", "Intresse (1–5)", "Mendelow-ruta",
    "RACI", "Kommunikationsbehov", "Anteckning",
]


def build_analys(wb, rows):
    ws = wb.create_sheet("Analys")
    ws.sheet_properties.tabColor = C_HEADER_MED
    widths = [4, 24, 15, 15, 18, 9, 30, 34]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    _banner(ws, "H", "📊 Analys — Mendelow & RACI")
    _table_header(ws, 2, ANALYS_HEADERS, bg=C_HEADER_DARK)

    data_start, data_end = 3, 2 + rows
    for idx in range(rows):
        r = data_start + idx
        bg = C_CREAM if idx % 2 == 0 else C_CREAM_DARK

        a = ws.cell(row=r, column=1, value=f'=IF({RAD}!B{r}<>"",ROW()-2,"")')
        a.font = Font(name="Arial", bold=True, size=9, color=C_INK_SOFT)
        a.fill = fill(bg)
        a.border = border_thin()
        a.alignment = Alignment(horizontal="center")

        b = ws.cell(row=r, column=2, value=f'=IF({RAD}!B{r}="","",{RAD}!B{r})')
        b.font = Font(name="Arial", size=9, color=C_INK)
        b.fill = fill(bg)
        b.border = border_thin()
        b.alignment = Alignment(vertical="center")

        for col in (3, 4):  # Inflytande, Intresse — inmatning
            c = ws.cell(row=r, column=col)
            c.font = Font(name="Arial", bold=True, size=10, color=C_INPUT_VAL)
            c.fill = fill(C_INPUT_BG)
            c.border = input_border()
            c.alignment = Alignment(horizontal="center")

        mend = (
            f'=IF(OR($C{r}="",$D{r}=""),"",'
            f'IF(AND($C{r}>=Settings!$C$8,$D{r}>=Settings!$C$9),"Hantera nära",'
            f'IF(AND($C{r}>=Settings!$C$8,$D{r}<Settings!$C$9),"Håll nöjd",'
            f'IF(AND($C{r}<Settings!$C$8,$D{r}>=Settings!$C$9),"Håll informerad",'
            f'"Bevaka"))))'
        )
        e = ws.cell(row=r, column=5, value=mend)
        e.font = Font(name="Arial", bold=True, size=9, color=C_FORMULA)
        e.fill = fill(bg)
        e.border = border_thin()
        e.alignment = Alignment(horizontal="center")

        for col in (6, 7, 8):  # RACI, Kommunikationsbehov, Anteckning — inmatning
            c = ws.cell(row=r, column=col)
            c.font = Font(name="Arial", size=9, color=C_INK)
            c.fill = fill(C_INPUT_BG)
            c.border = input_border()
            c.alignment = Alignment(vertical="center",
                                    wrap_text=(col in (7, 8)),
                                    horizontal="center" if col == 6 else "left")
        ws.row_dimensions[r].height = 18

    # Inmatningsvalidering
    dv_score = DataValidation(type="whole", operator="between", formula1="1",
                              formula2="5", allow_blank=True)
    dv_raci = DataValidation(type="list", formula1='"R,A,C,I,—"', allow_blank=True)
    ws.add_data_validation(dv_score)
    dv_score.add(f"C{data_start}:D{data_end}")
    ws.add_data_validation(dv_raci)
    dv_raci.add(f"F{data_start}:F{data_end}")

    # Färgkodning av Mendelow-rutan
    rng = f"E{data_start}:E{data_end}"
    rules = [
        ("Hantera nära", C_MANAGE),
        ("Håll nöjd", C_SATISFY),
        ("Håll informerad", C_INFORM),
        ("Bevaka", C_MONITOR),
    ]
    for label, color in rules:
        ws.conditional_formatting.add(
            rng,
            FormulaRule(formula=[f'$E{data_start}="{label}"'], fill=fill(color)),
        )

    _build_matrix(ws, data_start, data_end)
    ws.freeze_panes = "C3"


def _build_matrix(ws, data_start, data_end):
    """2×2 makt/intresse-matris som räknar intressenter per Mendelow-ruta."""
    rng = f"$E${data_start}:$E${data_end}"
    r0 = data_end + 2

    ws.merge_cells(f"B{r0}:E{r0}")
    ws[f"B{r0}"] = "MAKT/INTRESSE-MATRIS (Mendelow) — antal intressenter per ruta"
    ws[f"B{r0}"].font = Font(name="Arial", bold=True, size=9, color=C_WHITE)
    ws[f"B{r0}"].fill = fill(C_HEADER_MED)
    ws[f"B{r0}"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[r0].height = 20

    for col, txt in ((4, "Lågt intresse"), (5, "Högt intresse")):
        c = ws.cell(row=r0 + 1, column=col, value=txt)
        c.font = Font(name="Arial", bold=True, italic=True, size=9, color=C_INK_SOFT)
        c.fill = fill(C_CREAM_DARK)
        c.alignment = Alignment(horizontal="center")
        c.border = border_thin()

    cells = [
        (r0 + 2, "Högt inflytande", ("Håll nöjd", C_SATISFY), ("Hantera nära", C_MANAGE)),
        (r0 + 3, "Lågt inflytande", ("Bevaka", C_MONITOR), ("Håll informerad", C_INFORM)),
    ]
    for r, side_label, left, right in cells:
        s = ws.cell(row=r, column=3, value=side_label)
        s.font = Font(name="Arial", bold=True, italic=True, size=9, color=C_INK_SOFT)
        s.fill = fill(C_CREAM_DARK)
        s.alignment = Alignment(horizontal="right", vertical="center")
        s.border = border_thin()
        for col, (label, color) in ((4, left), (5, right)):
            c = ws.cell(
                row=r, column=col,
                value=f'=CONCATENATE("{label}: ",COUNTIF({rng},"{label}"))',
            )
            c.font = Font(name="Arial", bold=True, size=9, color=C_INK)
            c.fill = fill(color)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border_thin()
        ws.row_dimensions[r].height = 28

    total_r = r0 + 4
    ws[f"C{total_r}"] = "Kartlagda intressenter"
    ws[f"C{total_r}"].font = Font(name="Arial", bold=True, size=9, color=C_WHITE)
    ws[f"C{total_r}"].fill = fill(C_HEADER_DARK)
    ws[f"C{total_r}"].border = border_thin()
    ws[f"C{total_r}"].alignment = Alignment(horizontal="right")
    tot = ws.cell(
        row=total_r, column=4,
        value=f'=COUNTA($B${data_start}:$B${data_end})',
    )
    tot.font = Font(name="Arial", bold=True, size=9, color=C_GOLD_LIGHT)
    tot.fill = fill(C_HEADER_DARK)
    tot.border = border_thin()
    tot.alignment = Alignment(horizontal="center")


# ════════════════════════════════════════════════════════════
# FLIK 5 – File History
# ════════════════════════════════════════════════════════════
def build_history(wb):
    ws = wb.create_sheet("File History")
    ws.sheet_properties.tabColor = "8B6914"
    for col, w in zip("ABCD", (10, 16, 22, 56)):
        ws.column_dimensions[col].width = w

    _banner(ws, "D", "📋 Versionshistorik")
    _table_header(ws, 2, ["Version", "Datum", "Skapad av", "Ändringar"],
                  bg=C_HEADER_DARK)

    history = [
        ("v1.0", date.today().strftime("%Y-%m-%d"), "Clio / Fredrik Arvas",
         "Initial version. 5-fliksstruktur. Rådata från Slack-metadata, Analys med "
         "automatisk Mendelow-ruta (trösklar i Settings), RACI och makt/intresse-matris. "
         "AIAB-färgpalett."),
    ]
    for i, (ver, d, author, changes) in enumerate(history):
        r = 3 + i
        bg = C_CREAM if i % 2 == 0 else C_CREAM_DARK
        for col, val in enumerate([ver, d, author, changes], 1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = Font(name="Arial", size=9, color=C_INPUT_VAL if col == 1 else C_INK,
                          bold=(col == 1))
            c.fill = fill(bg)
            c.border = border_thin()
            c.alignment = Alignment(wrap_text=(col == 4), vertical="top")
        ws.row_dimensions[r].height = 46


# ── CSV ──────────────────────────────────────────────────────
def _read_csv(path):
    """Läser en CSV och returnerar rader som listor (hoppar ev. rubrikrad)."""
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as fh:
        for cells in csv.reader(fh):
            if not any(c.strip() for c in cells):
                continue
            if cells and cells[0].strip().lower() in ("namn", "name", "#"):
                continue
            rows.append([c.strip() for c in cells])
    return rows


def main():
    parser = argparse.ArgumentParser(
        description="Bygg Excel-mall för intressentanalys (Slack stakeholder mapping)."
    )
    parser.add_argument("--out", default="Clio_Intressentanalys_Mall_v1.0.xlsx",
                        help="Sökväg för output-filen.")
    parser.add_argument("--csv", help="Valfri CSV som förfyller Rådata-fliken.")
    parser.add_argument("--rows", type=int, default=40,
                        help="Antal intressentrader att förbereda (default 40).")
    parser.add_argument("--kund", default="[Kundnamn]",
                        help="Kund-/projektnamn till Settings.")
    args = parser.parse_args()

    csv_rows = _read_csv(args.csv) if args.csv else []
    rows = max(args.rows, len(csv_rows), 1)

    wb = Workbook()
    build_instructions(wb)
    build_settings(wb, args.kund)
    build_radata(wb, rows, csv_rows)
    build_analys(wb, rows)
    build_history(wb)

    out = Path(args.out).expanduser()
    wb.save(out)
    print(f"Sparad: {out.resolve()}")
    if csv_rows:
        print(f"Förfyllde {len(csv_rows)} rad(er) i Rådata från {args.csv}")


if __name__ == "__main__":
    main()
