"""word_builder.py — Konverterar .md-rapport till Word (.docx) med python-docx."""

from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path

logger = logging.getLogger(__name__)

FOOTER_URL = "https://fredrik.arvas.se/clio-research/"

# AIAB-fargpalett (RGBColor)
NAVY  = (42,  63,  111)
BLUE  = (74,  111, 165)
GOLD  = (200, 168, 75)
BROWN = (61,  46,  10)
CREAM = (247, 242, 232)
BEIGE = (237, 229, 208)
GREY  = (180, 180, 180)
WHITE = (255, 255, 255)


def build_docx(md_path: Path) -> Path:
    from docx import Document

    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    _setup_styles(doc)
    _setup_page(doc)
    _add_footer(doc)

    lines = text.splitlines()
    i = 0
    skip_toc_lines = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Byt ut markdown-TOC mot Word-falt med automatiska sidhanvisningar
        if stripped.startswith("## ") and "nneh" in stripped:
            doc.add_heading("Innehållsförteckning", level=2)
            _add_toc_field(doc)
            # Hoppa over alla efterfoljande TOC-rader tills nasta sektion
            i += 1
            while i < len(lines):
                s = lines[i].strip()
                if s.startswith("## ") or s == "---":
                    break
                i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_table(doc, table_lines)
            continue

        _render_line(doc, line)
        i += 1

    out_path = md_path.with_suffix(".docx")
    doc.save(str(out_path))
    logger.info("[word_builder] Docx sparad: %s", out_path)
    return out_path


def _add_toc_field(doc) -> None:
    """Infogar ett Word TOC-falt som populeras med sidnummer nar Word oppnar dokumentet."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)

    run = p.add_run()
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(*BROWN)

    # BEGIN
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    fld_begin.set(qn("w:dirty"), "true")
    run._r.append(fld_begin)

    # Instruktion: TOC over Heading 1-3, med hyperlank och sidhänvisning
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run._r.append(instr)

    # SEPARATE (platshallare-text som visas tills anvandaren uppdaterar)
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)

    # PLACEHOLDER-text
    placeholder = OxmlElement("w:t")
    placeholder.text = "[ Öppna i Word och tryck Ctrl+A sedan F9 för att uppdatera sidnummer ]"
    run._r.append(placeholder)

    # END
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def _setup_styles(doc):
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    styles = doc.styles

    # Normal
    n = styles["Normal"]
    n.font.name = "Arial"
    n.font.size = Pt(10)
    n.font.color.rgb = RGBColor(*BROWN)
    n.paragraph_format.space_after = Pt(4)

    # Heading 1 — titel
    h1 = styles["Heading 1"]
    h1.font.name = "Arial"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(*WHITE)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(12)

    # Heading 2 — sektionsrubrik
    h2 = styles["Heading 2"]
    h2.font.name = "Arial"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(*NAVY)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    # Heading 3
    h3 = styles["Heading 3"]
    h3.font.name = "Arial"
    h3.font.size = Pt(11)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(*BLUE)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(2)


def _setup_page(doc):
    from docx.shared import Cm
    section = doc.sections[0]
    section.page_width  = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)


def _add_footer(doc):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False

    para = footer.paragraphs[0]
    para.clear()

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = para.add_run("clio-research  |  " + FOOTER_URL + "  |  Sida ")
    run.font.name = "Arial"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(*GREY)

    # Sidnummer-falt
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.text = " PAGE "
    run._r.append(instr)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld2)

    run2 = para.add_run(" av ")
    run2.font.name = "Arial"
    run2.font.size = Pt(7)
    run2.font.color.rgb = RGBColor(*GREY)

    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "begin")
    run2._r.append(fld3)
    instr2 = OxmlElement("w:instrText")
    instr2.text = " NUMPAGES "
    run2._r.append(instr2)
    fld4 = OxmlElement("w:fldChar")
    fld4.set(qn("w:fldCharType"), "end")
    run2._r.append(fld4)

    run3 = para.add_run("  |  " + datetime.now().strftime("%Y-%m-%d"))
    run3.font.name = "Arial"
    run3.font.size = Pt(7)
    run3.font.color.rgb = RGBColor(*GREY)


def _add_table(doc, table_lines: list) -> None:
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    rows = []
    for line in table_lines:
        if re.match(r"^\|[-| :]+\|$", line):
            continue
        cells = [_clean(c.strip()) for c in line.strip("|").split("|")]
        if any(cells):
            rows.append(cells)

    if not rows:
        return

    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"

    # Kolumnbredder i cm (totalt ~16cm)
    col_w_presets = {
        2: [8.0, 8.0],
        3: [5.5, 5.5, 5.0],
        4: [4.0, 4.5, 4.0, 3.5],
        5: [2.5, 4.5, 2.0, 3.0, 4.0],
        6: [0.8, 5.5, 1.2, 2.0, 3.0, 3.5],
        7: [0.8, 4.5, 1.2, 1.8, 2.5, 2.7, 2.5],
    }
    col_widths = col_w_presets.get(n_cols, [round(16.0 / n_cols, 1)] * n_cols)

    for row_idx, row_data in enumerate(rows):
        is_header = (row_idx == 0)
        tr = table.rows[row_idx]
        for ci in range(n_cols):
            cell = table.cell(row_idx, ci)
            cell_text = row_data[ci] if ci < len(row_data) else ""

            # Bredd
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(int(col_widths[ci] * 567)))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)

            # Bakgrundsfarg
            if is_header:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "%02X%02X%02X" % NAVY)
                tcPr.append(shd)
            elif row_idx % 2 == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "%02X%02X%02X" % CREAM)
                tcPr.append(shd)
            else:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "%02X%02X%02X" % BEIGE)
                tcPr.append(shd)

            para = cell.paragraphs[0]
            para.clear()
            run = para.add_run(cell_text)
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.bold = is_header
            run.font.color.rgb = RGBColor(*(WHITE if is_header else BROWN))

    doc.add_paragraph()


def _render_line(doc, line: str) -> None:
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    stripped = line.strip()

    if stripped.startswith("# ") and not stripped.startswith("## "):
        p = doc.add_heading(_clean(stripped[2:]), level=1)
        # Navy bakgrund pa H1
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "%02X%02X%02X" % NAVY)
        pPr.append(shd)

    elif stripped.startswith("## "):
        doc.add_heading(_clean(stripped[3:]), level=2)

    elif stripped.startswith("### "):
        doc.add_heading(_clean(stripped[4:]), level=3)

    elif stripped == "---":
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "%02X%02X%02X" % GREY)
        pBdr.append(bottom)
        pPr.append(pBdr)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)

    elif stripped == "":
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)

    elif stripped.startswith("- ") or stripped.startswith("* "):
        p = doc.add_paragraph(style="List Bullet")
        _add_run(p, _clean(stripped[2:]))

    elif re.match(r"^\d+\.\s", stripped):
        # Kallforteckning-rader — liten text
        p = doc.add_paragraph()
        run = _add_run(p, _clean(stripped))
        run.font.size = Pt(8)

    else:
        p = doc.add_paragraph()
        _add_inline(p, stripped)


def _add_run(para, text: str, bold=False, italic=False, size=None):
    from docx.shared import Pt, RGBColor
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*BROWN)
    if size:
        run.font.size = Pt(size)
    return run


def _add_inline(para, text: str):
    """Hanterar **bold** och *italic* inline."""
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            _add_run(para, _clean(part[2:-2]), bold=True)
        elif part.startswith("*") and part.endswith("*"):
            _add_run(para, _clean(part[1:-1]), italic=True)
        else:
            _add_run(para, _clean(part))


def _clean(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text.strip()
