"""
Tester för ingest_minne.py — felhantering per fil, utan Qdrant eller OpenAI.
"""

from __future__ import annotations

import sys
from pathlib import Path
from unittest.mock import patch

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

import ingest_minne


def test_ingest_files_all_ok():
    files = [Path("a.txt"), Path("b.txt")]
    with patch.object(ingest_minne, "ingest_file", return_value=3) as mock_ingest:
        total, failed = ingest_minne.ingest_files(files, "mem_test")
    assert total == 6
    assert failed == []
    assert mock_ingest.call_count == 2


def test_ingest_files_continues_after_failure(capsys):
    files = [Path("a.txt"), Path("trasig.docx"), Path("c.txt")]

    def fake_ingest(path, collection, force=False):
        if path.name == "trasig.docx":
            raise ValueError("is not a Word file")
        return 2

    with patch.object(ingest_minne, "ingest_file", side_effect=fake_ingest):
        total, failed = ingest_minne.ingest_files(files, "mem_test")
    assert total == 4
    assert failed == [Path("trasig.docx")]
    assert "FEL" in capsys.readouterr().out


def test_main_exits_2_when_files_fail(tmp_path):
    (tmp_path / "x.txt").write_text("text")
    argv = ["ingest_minne.py", "--source", str(tmp_path), "--collection", "mem_test"]
    with patch.object(sys, "argv", argv), \
         patch.object(ingest_minne, "ensure_collection"), \
         patch.object(ingest_minne, "ingest_file", side_effect=RuntimeError("boom")):
        with pytest.raises(SystemExit) as exc:
            ingest_minne.main()
    assert exc.value.code == 2


# --- typigenkänning: ändelsen ljuger, innehållet avgör ----------------------

def _make_pdf(path: Path, text: str) -> Path:
    import pymupdf
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), text)
    doc.save(str(path))
    return path


def _make_docx(path: Path, text: str) -> Path:
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _make_pptx(path: Path, text: str, template: bool = False) -> Path:
    import zipfile
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = text
    tmp = path.with_name("_tmp.pptx")
    prs.save(str(tmp))
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if template and item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml",
                                    b"presentationml.template.main+xml")
            zout.writestr(item, data)
    tmp.unlink()
    return path


def _make_xlsx(path: Path, text: str) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Rubrik", text])
    wb.save(str(path))
    return path


@pytest.mark.parametrize("maker, name, kind", [
    (_make_pdf,  "egentligen_pdf.docx",  "pdf"),
    (_make_pdf,  "egentligen_pdf.pptx",  "pdf"),
    (_make_docx, "egentligen_docx.pptx", "docx"),
    (_make_pptx, "egentligen_pptx.docx", "pptx"),
    (_make_xlsx, "egentligen_xlsx.pptx", "xlsx"),
    (_make_docx, "riktig.docx",          "docx"),
])
def test_extract_text_follows_content_not_extension(tmp_path, maker, name, kind):
    path = maker(tmp_path / name, "Styrgruppsmote protokoll")
    assert ingest_minne.detect_kind(path) == kind
    assert "Styrgruppsmote protokoll" in ingest_minne.extract_text(path)


def test_pptx_template_disguised_as_docx_is_read(tmp_path):
    path = _make_pptx(tmp_path / "Projektdirektiv.docx", "Projektdirektiv", template=True)
    assert ingest_minne.detect_kind(path) == "pptx"
    assert ingest_minne.extract_text(path) == "[Bild 1] Projektdirektiv"


def test_txt_uses_extension(tmp_path):
    path = tmp_path / "anteckning.txt"
    path.write_text("vanlig text")
    assert ingest_minne.detect_kind(path) == "txt"
    assert ingest_minne.extract_text(path) == "vanlig text"


def test_unreadable_file_raises_and_is_reported(tmp_path):
    path = tmp_path / "trasig.docx"
    path.write_bytes(b"inte en zip och inte en pdf")
    assert ingest_minne.detect_kind(path) == "docx"
    total, failed = ingest_minne.ingest_files([path], "mem_test")
    assert (total, failed) == (0, [path])


def test_main_returns_normally_when_all_ok(tmp_path):
    (tmp_path / "x.txt").write_text("text")
    argv = ["ingest_minne.py", "--source", str(tmp_path), "--collection", "mem_test"]
    with patch.object(sys, "argv", argv), \
         patch.object(ingest_minne, "ensure_collection"), \
         patch.object(ingest_minne, "ingest_file", return_value=1):
        ingest_minne.main()
