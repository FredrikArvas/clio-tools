"""
output.py — Genererar Word-filer från förbättrings-JSON.

Fyra filer per pass:
  INPUT  (från pairs i JSON):
    {stem}_original_sv.docx   — svensk källtext
    {stem}_original_en.docx   — engelsk originaltext

  OUTPUT (från improvements i JSON):
    {stem}_proposed_en.docx   — förbättrad engelska (för Words Compare Documents)
    {stem}_commented.docx     — ORIGINAL/PROPOSED/WHY i färg, för granskning

Användning:
  python output.py --improvements testdata/backcover_improvements.json
  python output.py --improvements testdata/chapter_01_improvements.json --out-dir output/
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def _require_docx():
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        return Document, Pt, RGBColor
    except ImportError:
        print("python-docx saknas — installera: pip install python-docx", file=sys.stderr)
        sys.exit(1)


def _add_label(para, label: str, Pt, RGBColor):
    run = para.add_run(f"[{label}]  ")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def _is_placeholder(text: str) -> bool:
    return text.startswith("[")


def build_originals(data: dict, stem: str, out_dir: Path) -> tuple[Path, Path]:
    """Generera original_sv.docx och original_en.docx från pairs."""
    Document, Pt, RGBColor = _require_docx()
    all_pairs = data.get("all_pairs", [])
    title_en = data.get("title_en", "")
    title_sv = data.get("title_sv", "")

    paths = []
    for lang, title, key in [("sv", title_sv, "sv"), ("en", title_en, "en")]:
        doc = Document()
        doc.add_heading(title or f"Original ({lang.upper()})", level=1)
        doc.add_paragraph()
        for pair in all_pairs:
            text = pair.get(key, pair.get("en", ""))
            if text and not _is_placeholder(text):
                doc.add_paragraph(text)
        path = out_dir / f"{stem}_original_{lang}.docx"
        doc.save(path)
        paths.append(path)
        print(f"Sparat: {path}")

    return paths[0], paths[1]


def build_proposed(data: dict, stem: str, out_dir: Path) -> Path:
    """Generera proposed_en.docx — förbättrad engelska för Words Compare."""
    Document, Pt, RGBColor = _require_docx()
    all_pairs = data.get("all_pairs", [])
    improvements_by_id = {item["id"]: item for item in data.get("improvements", [])}
    title_en = data.get("title_en", "")

    doc = Document()
    doc.add_heading(f"{title_en} — Proposed EN", level=1)
    doc.add_paragraph()

    for pair in all_pairs:
        pid = pair["id"]
        item = improvements_by_id.get(pid)
        if item:
            doc.add_paragraph(item["en_proposed"])
        else:
            text = pair.get("en", "")
            if text and not _is_placeholder(text):
                doc.add_paragraph(text)

    path = out_dir / f"{stem}_proposed_en.docx"
    doc.save(path)
    print(f"Sparat: {path}")
    return path


def build_commented(data: dict, stem: str, out_dir: Path) -> Path:
    """Generera commented.docx — ORIGINAL/PROPOSED/WHY i färg."""
    Document, Pt, RGBColor = _require_docx()
    all_pairs = data.get("all_pairs", [])
    improved_ids = {item["id"] for item in data.get("improvements", [])}
    improvements_by_id = {item["id"]: item for item in data.get("improvements", [])}
    title_en = data.get("title_en", "")

    doc = Document()
    doc.add_heading(f"{title_en} — Commented", level=1)
    doc.add_paragraph(
        f"Analysis: {data.get('analysis_date', '')}  |  "
        f"{len(improved_ids)} improvements of {len(all_pairs)} paragraphs"
    ).runs[0].font.size = Pt(9)
    doc.add_paragraph()

    for pair in all_pairs:
        pid = pair["id"]

        if pid not in improved_ids:
            p = doc.add_paragraph()
            _add_label(p, "UNCHANGED", Pt, RGBColor)
            en_text = pair.get("en", "")
            if not _is_placeholder(en_text):
                run = p.add_run(en_text)
                run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
                run.font.size = Pt(10)
        else:
            item = improvements_by_id[pid]

            p_orig = doc.add_paragraph()
            _add_label(p_orig, "ORIGINAL EN", Pt, RGBColor)
            en_text = item["en_original"]
            if not _is_placeholder(en_text):
                run = p_orig.add_run(en_text)
                run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
                run.font.strike = True
                run.font.size = Pt(11)
            else:
                run = p_orig.add_run("(missing from English original)")
                run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
                run.italic = True
                run.font.size = Pt(10)

            p_sv = doc.add_paragraph()
            _add_label(p_sv, "SWEDISH", Pt, RGBColor)
            run_sv = p_sv.add_run(item["sv_improved"])
            run_sv.font.color.rgb = RGBColor(0x66, 0x44, 0x00)
            run_sv.italic = True
            run_sv.font.size = Pt(10)

            p_prop = doc.add_paragraph()
            _add_label(p_prop, "PROPOSED EN", Pt, RGBColor)
            run_prop = p_prop.add_run(item["en_proposed"])
            run_prop.font.color.rgb = RGBColor(0x1A, 0x7A, 0x1A)
            run_prop.bold = True
            run_prop.font.size = Pt(11)

            p_why = doc.add_paragraph()
            _add_label(p_why, "WHY", Pt, RGBColor)
            run_why = p_why.add_run(item["motivation"])
            run_why.font.color.rgb = RGBColor(0x1A, 0x4A, 0x99)
            run_why.italic = True
            run_why.font.size = Pt(9)

        doc.add_paragraph()

    path = out_dir / f"{stem}_commented.docx"
    doc.save(path)
    print(f"Sparat: {path}")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description="clio-bookrev output — improvements JSON -> 4 Word-filer")
    parser.add_argument("--improvements", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, help="Utdatakatalog (default: output/)")
    args = parser.parse_args()

    if not args.improvements.exists():
        print(f"Fil saknas: {args.improvements}", file=sys.stderr)
        sys.exit(1)

    data = json.loads(args.improvements.read_text(encoding="utf-8"))
    stem = args.improvements.stem.replace("_improvements", "")
    out_dir = args.out_dir or args.improvements.parent.parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"\n-- Genererar 4 filer for: {stem} --")
    build_originals(data, stem, out_dir)
    build_proposed(data, stem, out_dir)
    build_commented(data, stem, out_dir)
    print(f"\nKlar. Word Compare: original_en vs proposed_en")


if __name__ == "__main__":
    main()
